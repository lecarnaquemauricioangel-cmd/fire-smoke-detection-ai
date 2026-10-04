"""
Generación de reportes en Excel (.xlsx) a partir del historial de eventos
almacenado en SQLite, usando openpyxl.
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database

BASE_DIR = Path(__file__).resolve().parent
REPORTES_DIR = BASE_DIR / "static" / "reportes"
REPORTES_DIR.mkdir(parents=True, exist_ok=True)

COLOR_SEVERIDAD = {
    "ALTA": "FFC7CE",
    "MEDIA": "FFEB9C",
    "BAJA": "C6EFCE",
}


def generar_excel(nombre_archivo: str = "reporte_eventos.xlsx") -> str:
    """
    Exporta el historial completo de eventos a un archivo .xlsx.

    Retorna la ruta absoluta (str) del archivo generado.
    """
    eventos = database.obtener_todos()

    wb = Workbook()
    ws = wb.active
    ws.title = "Historial de Eventos"

    encabezados = ["ID", "Fecha", "Fuente", "Clase", "Confianza", "Detalle"]
    ws.append(encabezados)

    encabezado_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    encabezado_font = Font(color="FFFFFF", bold=True)
    borde_fino = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    for celda in ws[1]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde_fino

    for evento in eventos:
        fila = [
            evento["id"],
            evento["fecha"],
            evento["fuente"],
            evento["clase"],
            round(evento["confianza"], 3),
            evento["detalle"],
        ]
        ws.append(fila)

        fila_actual = ws.max_row
        detalle_texto = (evento["detalle"] or "").upper()
        color = None
        for nivel, hexc in COLOR_SEVERIDAD.items():
            if nivel in detalle_texto:
                color = hexc
                break

        for celda in ws[fila_actual]:
            celda.border = borde_fino
            celda.alignment = Alignment(horizontal="center", vertical="center")
            if color:
                celda.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    anchos = [8, 20, 28, 12, 12, 45]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.freeze_panes = "A2"

    ruta_salida = REPORTES_DIR / nombre_archivo
    wb.save(ruta_salida)
    return str(ruta_salida)
